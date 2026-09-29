from pathlib import Path

import pytest

from eka.ingestion.parsers import MIME_BY_EXT, UnsupportedFile, drop_repeated_lines, parse, sniff_mime


def mime(name: str) -> str:
    return MIME_BY_EXT[Path(name).suffix]


def kinds(blocks) -> list[tuple[str, str]]:  # type: ignore[no-untyped-def]
    return [(b.kind, b.text) for b in blocks]


def test_docx(tmp_path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("Leave Policy", level=1)
    d.add_paragraph("Staff  get\u00a010 days.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Type", "Days", "Sick", "10"
    d.add_paragraph("After the table.")
    p = tmp_path / "a.docx"
    d.save(str(p))
    blocks = parse(p, mime("a.docx"))
    assert [b.kind for b in blocks] == ["heading", "paragraph", "table", "paragraph"]
    assert blocks[0].level == 1 and blocks[1].text == "Staff get 10 days."
    assert blocks[2].rows == [["Type", "Days"], ["Sick", "10"]]


def test_xlsx_header_is_first_non_empty_row(tmp_path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Codes"
    ws.append([None, None])
    ws.append(["Code", "Meaning"])
    ws.append(["ERR-4032", "Quota exceeded"])
    p = tmp_path / "a.xlsx"
    wb.save(p)
    blocks = parse(p, mime("a.xlsx"))
    assert blocks[0].text == "Codes"
    assert blocks[1].rows == [["Code", "Meaning"], ["ERR-4032", "Quota exceeded"]]


def test_pptx_titles_and_notes(tmp_path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Roadmap"
    s.placeholders[1].text = "Ship v0.1 in December."
    s.notes_slide.notes_text_frame.text = "Mention the demo."
    p = tmp_path / "a.pptx"
    prs.save(p)
    blocks = parse(p, mime("a.pptx"))
    assert kinds(blocks) == [
        ("heading", "Roadmap"),
        ("paragraph", "Ship v0.1 in December."),
        ("paragraph", "Mention the demo."),
    ]
    assert all(b.page == 1 for b in blocks)


def test_html_drops_chrome_and_keeps_order(tmp_path: Path) -> None:
    p = tmp_path / "a.html"
    p.write_text(
        "<nav>Menu</nav><h2>Title</h2><p>Body <b>text</b></p><ul><li>one</li></ul>"
        "<table><tr><th>A</th></tr><tr><td>1</td></tr></table><script>x()</script><footer>f</footer>"
    )
    blocks = parse(p, mime("a.html"))
    assert kinds(blocks)[:3] == [("heading", "Title"), ("paragraph", "Body text"), ("list", "one")]
    assert blocks[3].rows == [["A"], ["1"]]
    assert "Menu" not in str(blocks) and "x()" not in str(blocks)


def test_markdown(tmp_path: Path) -> None:
    p = tmp_path / "a.md"
    p.write_text("# Guide\n\nIntro.\n\n```\ncode  line\n```\n")
    assert kinds(parse(p, mime("a.md"))) == [("heading", "Guide"), ("paragraph", "Intro."), ("code", "code line")]


def _minimal_pdf(pages: list[str]) -> bytes:
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", ""]
    kids = []
    for text in pages:
        n = len(objs) + 1
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {n + 1} 0 R "
            f"/Resources << /Font << /F1 {n + 2} 0 R >> >> >>"
        )
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
        kids.append(f"{n} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>"
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def test_pdf(tmp_path: Path) -> None:
    p = tmp_path / "a.pdf"
    p.write_bytes(_minimal_pdf(["First page text.", "Second page text."]))
    blocks = parse(p, mime("a.pdf"))
    assert [(b.text, b.page) for b in blocks] == [("First page text.", 1), ("Second page text.", 2)]


def test_repeated_header_lines_dropped() -> None:
    pages = [f"ACME Confidential\nContent {i}\nPage {i} of 4" for i in range(1, 5)]
    assert drop_repeated_lines(pages) == [f"Content {i}" for i in range(1, 5)]


def test_sniff_rejects_renamed_files() -> None:
    assert sniff_mime("x.pdf", b"%PDF-1.7") == "application/pdf"
    assert sniff_mime("x.docx", b"PK\x03\x04abcd").endswith("document")
    with pytest.raises(UnsupportedFile):
        sniff_mime("x.pdf", b"MZ\x90\x00")
    with pytest.raises(UnsupportedFile):
        sniff_mime("x.exe", b"MZ")
