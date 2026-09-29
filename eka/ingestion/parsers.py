"""File bytes -> ordered blocks (D8). One function per format; parse() picks by MIME type."""

import re
import unicodedata
from collections import Counter
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path

from bs4 import BeautifulSoup, Tag

from eka.core.config import settings

MIME_BY_EXT = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".html": "text/html",
    ".htm": "text/html",
    ".md": "text/markdown",
    ".txt": "text/plain",
}
XLSX_MAX_ROWS = 2000
OCR_MIN_CHARS = 20


@dataclass
class Block:
    kind: str  # heading | paragraph | table | code | list
    text: str = ""
    page: int | None = None
    level: int = 0  # heading level, 1 = top
    rows: list[list[str]] = field(default_factory=list)  # table: rows[0] is the header


class UnsupportedFile(ValueError):
    pass


def sniff_mime(filename: str, head: bytes) -> str:
    """MIME from the extension, confirmed against the leading bytes so a renamed file is rejected."""
    mime = MIME_BY_EXT.get(Path(filename).suffix.lower())
    if mime is None:
        raise UnsupportedFile(f"unsupported file type: {filename}")
    if mime == "application/pdf":
        ok = head.startswith(b"%PDF")
    elif mime.startswith("application/vnd.openxmlformats"):
        ok = head.startswith(b"PK\x03\x04")
    else:
        ok = b"\x00" not in head
    if not ok:
        raise UnsupportedFile(f"content does not match extension: {filename}")
    return mime


def normalise(text: str, keep_newlines: bool = False) -> str:
    text = unicodedata.normalize("NFKC", text)
    if keep_newlines:
        return "\n".join(re.sub(r"[ \t]+", " ", line).rstrip() for line in text.splitlines()).strip("\n")
    return re.sub(r"\s+", " ", text).strip()


def _cell(value: object) -> str:
    return normalise("" if value is None else str(value))


def _table(rows: Sequence[Sequence[object]], page: int | None = None) -> Block | None:
    clean = [[_cell(c) for c in r] for r in rows]
    clean = [r for r in clean if any(r)]
    return Block("table", page=page, rows=clean) if clean else None


def _paragraphs(text: str, page: int | None = None) -> list[Block]:
    return [Block("paragraph", normalise(p), page) for p in re.split(r"\n\s*\n", text) if p.strip()]


def _parse_pdf(path: Path) -> list[Block]:
    import pdfplumber

    texts: list[str] = []
    tables: list[list[Block]] = []
    with pdfplumber.open(path) as pdf:
        for n, page in enumerate(pdf.pages, start=1):
            page_tables: list[Block] = []
            rest = page
            for t in page.find_tables():
                rest = rest.outside_bbox(t.bbox)
                if b := _table(t.extract(), n):
                    page_tables.append(b)
            text = rest.extract_text() or ""
            if len(text.strip()) < OCR_MIN_CHARS and settings.ocr_enabled:
                import pytesseract

                text = pytesseract.image_to_string(page.to_image(resolution=300).original)
            texts.append(text)
            tables.append(page_tables)
    texts = drop_repeated_lines(texts)
    return [b for n, (text, tbl) in enumerate(zip(texts, tables, strict=True), 1) for b in _paragraphs(text, n) + tbl]


_PAGE_NUMBER = re.compile(r"^(page\s*)?\d+(\s*(of|/)\s*\d+)?$", re.IGNORECASE)


def drop_repeated_lines(pages: list[str]) -> list[str]:
    """Drops running headers and footers: lines on more than half the pages, and bare page numbers."""
    if len(pages) < 3:
        return pages
    seen: Counter[str] = Counter()
    for p in pages:
        seen.update({normalise(ln) for ln in p.splitlines() if ln.strip()})
    repeated = {k for k, c in seen.items() if c > len(pages) / 2}

    def keep(line: str) -> bool:
        n = normalise(line)
        return n not in repeated and not _PAGE_NUMBER.match(n)

    return ["\n".join(ln for ln in p.splitlines() if keep(ln)) for p in pages]


def _parse_docx(path: Path) -> list[Block]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    blocks: list[Block] = []
    for el in d.element.body.iterchildren():
        if el.tag.endswith("}tbl"):
            t = Table(el, d)
            if b := _table([[c.text for c in r.cells] for r in t.rows]):
                blocks.append(b)
        elif el.tag.endswith("}p"):
            p = Paragraph(el, d)
            text = normalise(p.text)
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style == "Title":
                blocks.append(Block("heading", text, level=1))
            elif m := re.match(r"Heading (\d)", style):
                blocks.append(Block("heading", text, level=int(m.group(1))))
            elif style.startswith("List"):
                blocks.append(Block("list", text))
            else:
                blocks.append(Block("paragraph", text))
    return blocks


def _parse_xlsx(path: Path) -> list[Block]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    blocks: list[Block] = []
    try:
        for ws in wb.worksheets:
            rows: list[list[object]] = []
            for row in ws.iter_rows(values_only=True):
                if not rows and not any(c is not None and str(c).strip() for c in row):
                    continue  # header = first non-empty row
                rows.append(list(row))
                if len(rows) > XLSX_MAX_ROWS:
                    break
            if b := _table(rows):
                blocks += [Block("heading", normalise(ws.title), level=1), b]
    finally:
        wb.close()
    return blocks


def _parse_pptx(path: Path) -> list[Block]:
    from pptx import Presentation

    blocks: list[Block] = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        title = slide.shapes.title
        if title is not None and title.text.strip():
            blocks.append(Block("heading", normalise(title.text), n, level=1))
        for shape in slide.shapes:
            if shape == title:
                continue
            if shape.has_table:
                if b := _table([[c.text for c in r.cells] for r in shape.table.rows], n):
                    blocks.append(b)
            elif shape.has_text_frame and shape.text_frame.text.strip():
                blocks += _paragraphs(shape.text_frame.text, n)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            blocks += _paragraphs(slide.notes_slide.notes_text_frame.text, n)
    return blocks


_HTML_BLOCKS = ("h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "table", "blockquote")


def _parse_html_text(html: str) -> list[Block]:
    soup = BeautifulSoup(html, "html.parser")
    for el in soup(["nav", "footer", "script", "style", "noscript", "header"]):
        el.decompose()
    blocks: list[Block] = []
    for el in soup.find_all(_HTML_BLOCKS):
        assert isinstance(el, Tag)
        if el.find_parent(_HTML_BLOCKS):
            continue  # nested inside a block already emitted
        if el.name == "table":
            rows = [[c.get_text(" ") for c in tr.find_all(["th", "td"])] for tr in el.find_all("tr")]
            if b := _table(rows):  # type: ignore[arg-type]
                blocks.append(b)
        elif el.name == "pre":
            blocks.append(Block("code", normalise(el.get_text(), keep_newlines=True)))
        elif (text := normalise(el.get_text(" "))) == "":
            continue
        elif el.name[0] == "h":
            blocks.append(Block("heading", text, level=int(el.name[1])))
        else:
            blocks.append(Block("list" if el.name == "li" else "paragraph", text))
    return blocks


def _read_text(path: Path) -> str:
    return path.read_bytes().decode("utf-8", errors="replace")


def parse(path: str | Path, mime: str) -> list[Block]:
    path = Path(path)
    if mime == "application/pdf":
        return _parse_pdf(path)
    if mime.endswith("wordprocessingml.document"):
        return _parse_docx(path)
    if mime.endswith("spreadsheetml.sheet"):
        return _parse_xlsx(path)
    if mime.endswith("presentationml.presentation"):
        return _parse_pptx(path)
    if mime == "text/html":
        return _parse_html_text(_read_text(path))
    if mime == "text/markdown":
        import markdown

        return _parse_html_text(markdown.markdown(_read_text(path), extensions=["tables", "fenced_code"]))
    if mime == "text/plain":
        return _paragraphs(_read_text(path))
    raise UnsupportedFile(mime)
